#!/usr/bin/env python3
"""
validar_jev.py — validação cega da classificação de eventos do JEV (amostra de 50 documentos)
=============================================================================================
O JEV vê só o título. O gabarito vem da leitura do documento inteiro (PDF da CVM).

    python validar_jev.py amostra    # sorteia a amostra, baixa os PDFs e extrai o texto (validacao/textos/)
    python validar_jev.py planilha   # gera validacao/jev_validacao_50.xlsx com a sugestão de gabarito
    python validar_jev.py comparar   # compara a planilha preenchida com as respostas do JEV

A resposta do JEV fica em validacao/jev_respostas_amostra.json e não aparece na planilha.
"""

from __future__ import annotations

import json
import random
import re
import sys
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path

_ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(_ROOT))
import jev_eventos as jev  # noqa: E402

PASTA = _ROOT / "validacao"
AMOSTRA = PASTA / "amostra_jev.json"
RESPOSTAS = PASTA / "jev_respostas_amostra.json"
SUGESTOES = PASTA / "gabarito_sugerido.json"
PLANILHA = PASTA / "jev_validacao_50.xlsx"
TEXTOS = PASTA / "textos"
SEMENTE = 20261007
# títulos classificados como "outro" que sugerem um evento: medem o que o JEV deixa passar
SUSPEITOS = re.compile(r"auditor|ressalva|republica|reapresenta.{0,30}(demonstra|dfp|itr)|retifica|recupera[cç][aã]o judicial|"
                       r"extrajudicial|waiver|inadimpl|vencimento antecipado|renuncia|renúncia|destitui|investiga|"
                       r"inqu[eé]rito|processo administrativo|impairment|moeda funcional", re.I)


def _texto_pdf(conteudo: bytes) -> str:
    """Texto do PDF com pypdf (puro Python)."""
    import io
    from pypdf import PdfReader
    return chr(10).join((pg.extract_text() or "") for pg in PdfReader(io.BytesIO(conteudo)).pages)


def _docs_e_cache():
    dados = json.loads(jev._DADOS.read_text(encoding="utf-8"))
    empresas = {e["cnpj"]: e for e in dados["empresas"] if e.get("mscore") is not None}
    docs = {d["id"]: d for d in jev.documentos(empresas)}
    cache = json.loads(jev._CACHE.read_text(encoding="utf-8"))
    return empresas, docs, cache


def amostra() -> None:
    empresas, docs, cache = _docs_e_cache()
    rnd = random.Random(SEMENTE)
    por_evento = defaultdict(list)
    for i, d in docs.items():
        r = cache.get(i)
        if not r:
            continue
        ev = r["evento"] if r["evento"] == "outro" or r["confianca"] >= jev.CONFIANCA_MINIMA else "baixa_confianca"
        por_evento[ev].append(i)
    alvo = {"auditoria": 99, "politica_contabil": 99, "investigacao": 99, "reestruturacao_divida": 12, "saida_executivo": 10}
    escolhidos = []
    for ev, n in alvo.items():
        ids = sorted(por_evento.get(ev, []))
        escolhidos += ids if n >= len(ids) else rnd.sample(ids, n)
    outros = sorted(por_evento["outro"])
    suspeitos = [i for i in outros if SUSPEITOS.search(docs[i]["titulo"])]
    faltam = 50 - len(escolhidos)
    n_susp = min(len(suspeitos), (faltam + 1) // 2)
    escolhidos += rnd.sample(suspeitos, n_susp)
    restantes = [i for i in outros if i not in escolhidos]
    escolhidos += rnd.sample(restantes, 50 - len(escolhidos))
    rnd.shuffle(escolhidos)  # a ordem da planilha não revela o grupo

    TEXTOS.mkdir(parents=True, exist_ok=True)
    itens, respostas = [], {}
    for n, i in enumerate(escolhidos, 1):
        d = docs[i]
        itens.append({"n": n, "id": i, "cnpj": d["cnpj"], "empresa": empresas[d["cnpj"]]["nome"], "data": d["data"],
                      "categoria": d["categoria"], "titulo": d["titulo"], "url": d["url"]})
        respostas[i] = cache[i]
        destino = TEXTOS / f"{n:02d}.txt"
        if destino.exists():
            continue
        try:
            req = urllib.request.Request(d["url"], headers={"User-Agent": "Mozilla/5.0 (validacao JEV)"})
            pdf = urllib.request.urlopen(req, timeout=120).read()
            destino.write_text(_texto_pdf(pdf), encoding="utf-8")
        except Exception as exc:
            destino.write_text(f"[não foi possível ler o documento: {exc}]", encoding="utf-8")
    PASTA.mkdir(parents=True, exist_ok=True)
    AMOSTRA.write_text(json.dumps(itens, ensure_ascii=False, indent=1), encoding="utf-8")
    RESPOSTAS.write_text(json.dumps(respostas, ensure_ascii=False, indent=1), encoding="utf-8")
    grupos = Counter("suspeito" if i in suspeitos else (cache[i]["evento"] if cache[i]["confianca"] >= jev.CONFIANCA_MINIMA else "baixa") for i in escolhidos)
    print(f"{len(itens)} documentos; grupos: {dict(grupos)}; textos em {TEXTOS.relative_to(_ROOT)}")


def planilha() -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    itens = json.loads(AMOSTRA.read_text(encoding="utf-8"))
    sug = json.loads(SUGESTOES.read_text(encoding="utf-8")) if SUGESTOES.exists() else {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Validar"
    ws.append(["nº", "Empresa", "Data", "Documento", "Título (link para o original)", "Sugestão pela leitura do documento",
               "Por quê", "Evento (sua decisão)", "Comentário"])
    for it in itens:
        s = sug.get(str(it["n"]), {})
        ws.append([it["n"], it["empresa"].title(), it["data"], it["categoria"], it["titulo"], s.get("evento", ""),
                   s.get("motivo", ""), s.get("evento", ""), ""])
        ws.cell(row=ws.max_row, column=5).hyperlink = it["url"]
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1C5CAB")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate([5, 30, 11, 18, 60, 22, 60, 22, 30]):
        ws.column_dimensions[chr(65 + i)].width = w
    for linha in ws.iter_rows(min_row=2):
        for c in linha:
            c.alignment = Alignment(vertical="top", wrap_text=True)
        linha[7].fill = PatternFill("solid", fgColor="FFF7D6")
    dv = DataValidation(type="list", formula1='"' + ",".join(jev.EVENTOS) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"H2:H{len(itens) + 1}")
    ws.freeze_panes = "B2"
    leg = wb.create_sheet("Critérios")
    leg.append(["Evento", "O que conta (versão 2 dos critérios)"])
    for k, v in jev.EVENTOS.items():
        leg.append([k, v])
    leg.append([])
    leg.append(["Como usar", "A coluna H já vem com a sugestão. Confirme ou troque pela sua leitura; a resposta do JEV não está nesta planilha."])
    leg.column_dimensions["A"].width = 24
    leg.column_dimensions["B"].width = 120
    for c in leg[1]:
        c.font = Font(bold=True)
    PLANILHA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(PLANILHA)
    print(f"planilha → {PLANILHA.relative_to(_ROOT)}")


def comparar() -> None:
    from openpyxl import load_workbook
    itens = {it["n"]: it for it in json.loads(AMOSTRA.read_text(encoding="utf-8"))}
    resp = json.loads(RESPOSTAS.read_text(encoding="utf-8"))
    ws = load_workbook(PLANILHA)["Validar"]
    pares = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None or not row[7]:
            continue
        r = resp[itens[int(row[0])]["id"]]
        jev_ev = r["evento"] if r["confianca"] >= jev.CONFIANCA_MINIMA else "outro"  # o site só mostra confiança ≥ 0,8
        pares.append((str(row[7]).strip(), jev_ev, int(row[0]), itens[int(row[0])]["titulo"]))
    acertos = sum(1 for g, j, *_ in pares if g == j)
    print(f"{len(pares)} documentos; concordância geral: {acertos}/{len(pares)} = {acertos / len(pares):.0%}")
    eventos = [k for k in jev.EVENTOS if k != "outro"]
    marcados = [(g, j) for g, j, *_ in pares if j != "outro"]
    reais = [(g, j) for g, j, *_ in pares if g != "outro"]
    if marcados:
        print(f"precisão (quando o site mostra um evento, ele está certo): {sum(g == j for g, j in marcados)}/{len(marcados)}")
    if reais:
        print(f"cobertura na amostra (eventos reais que o site mostra): {sum(g == j for g, j in reais)}/{len(reais)}")
    for ev in eventos:
        m = [(g, j) for g, j, *_ in pares if j == ev]
        if m:
            print(f"  {ev:24} {sum(g == j for g, j in m)}/{len(m)} certos")
    print("divergências:")
    for g, j, n, t in pares:
        if g != j:
            print(f"  nº {n:2}: gabarito {g}, JEV {j} | {t[:80]}")


if __name__ == "__main__":
    {"amostra": amostra, "planilha": planilha, "comparar": comparar}[sys.argv[1]]()
